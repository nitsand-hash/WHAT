#!/usr/bin/env python3
"""Convert a board export (xlsx) into src/api/data.json for the local BoardSDK.

Usage: python3 scripts/import-xlsx.py path/to/export.xlsx   (needs `pip install openpyxl`)

Only the fields the app actually shows are copied. Credentials and internal links in the export
(Login Email/Pass, WorkSpace ID, WS Email, HubSpot, Contract, Whitelist) are deliberately left out.
Rows under "Churn" and "Not yet signed" are skipped; only the "Active" group is imported.
"""
import json
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

SRC = Path(sys.argv[1])
OUT = Path(__file__).resolve().parent.parent / 'src' / 'api' / 'data.json'


def split(v):
    return [p.strip() for p in str(v).split(',') if p.strip()] if v else []


def iso(v):
    return v.strftime('%Y-%m-%dT%H:%M:%S') if isinstance(v, datetime) else None


def num(v):
    return v if isinstance(v, (int, float)) else None


def link(v):
    url = str(v).split(',')[0].strip() if v else ''
    return {'url': url, 'label': ''} if url.startswith('http') else None


wb = openpyxl.load_workbook(SRC, data_only=True)
rows = list(wb.worksheets[0].iter_rows(values_only=True))
header = list(rows[2])
col = {h: i for i, h in enumerate(header)}
get = lambda r, name: r[col[name]] if name in col else None

owners = {}
clients = []
group = None
for r in rows[3:]:
    if all(v is None for v in r) or r[0] == 'Subitems':
        continue
    if r[0] and all(v is None for v in r[1:]):
        group = r[0]  # group title row ("Churn", "Not yet signed"); the first group is untitled = Active
        continue
    if r[0] is None or group is not None or get(r, 'Client Status') != 'Active':
        continue  # subitem rows, churned and unsigned deals
    cs = split(get(r, 'CS owner'))
    cs_owner = []
    for name in cs:
        owners.setdefault(name, 100 + len(owners) + 1)
        cs_owner.append({'id': owners[name], 'name': name})
    clients.append({
        'id': str(get(r, 'Item ID (auto generated)')),
        'name': r[0],
        'clientStatus': get(r, 'Client Status'),
        'industry': get(r, 'Industry'),
        'health': get(r, 'Health'),
        'commsFlag': get(r, 'Comms Flag'),
        'lifecycleStage': get(r, 'Lifecycle Stage'),
        'csOwner': cs_owner,
        'weeklyStatus': get(r, 'Weekly Status') or '',
        'lastTouchpoint': iso(get(r, 'Last Touchpoint')),
        'lastTouchType': get(r, 'Last Touch Type'),
        'nextMonthlyCall': iso(get(r, 'Next Monthly Call')),
        'renewalDate': iso(get(r, 'Renewal Date')),
        'peakEvent': get(r, 'Peak Event') or '',
        'peakEventDate': iso(get(r, 'Peak Event Date')),
        'arr': num(get(r, 'ARR')),
        'impersCap': num(get(r, 'Impers. Cap')) or 0,
        'commCap': num(get(r, 'Comm. Cap')) or 0,
        'pocName': get(r, 'POC Name'),
        'pocEmail': get(r, 'POC Email'),
        'activeSocials': split(get(r, 'Active socials')),
        'signedProducts': split(get(r, 'Signed products')),
        'loa': link(get(r, 'LOA')),
        'trademarkDoc': link(get(r, 'Trademark')),
        'updatedAt': iso(get(r, 'Last Touchpoint')),
    })

by_id = {c['id'] for c in clients}
updates = []
for n, r in enumerate(list(wb['updates'].iter_rows(values_only=True))[2:]):
    item_id, _, _, _, user, created, body = r[:7]
    if str(item_id) not in by_id or not body:
        continue
    at = datetime.strptime(' '.join(str(created).split()), '%d/%B/%Y %I:%M:%S %p')
    updates.append({
        'id': str(r[9] or f'u{n}'),
        'itemId': str(item_id),
        'created_at': iso(at),
        'text_body': str(body),
        'creator': {'name': user or 'Unknown'},
    })

OUT.write_text(json.dumps({'clients': clients, 'updates': updates}, ensure_ascii=False, indent=1))
print(f'{len(clients)} clients, {len(updates)} updates -> {OUT}')
